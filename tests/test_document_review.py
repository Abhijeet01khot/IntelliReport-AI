
from io import BytesIO

from docx import Document

from utils.document_review import (
    extract_document_text,
    make_agent_result,
    run_agent_with_status,
)


def make_docx_bytes(text="A sample document for review."):
    document = Document()
    document.add_paragraph(text)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_extract_docx_text():
    content = extract_document_text(
        "sample.docx",
        make_docx_bytes(),
    )
    assert "sample document for review" in content


def test_extract_pdf_text():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)

    buffer = BytesIO()
    writer.write(buffer)

    # A blank PDF has no extractable text and should be rejected.
    try:
        extract_document_text("blank.pdf", buffer.getvalue())
    except ValueError as error:
        assert "readable text" in str(error).lower()
    else:
        raise AssertionError("Blank PDF should be rejected")


def test_reject_unsupported_file_type():
    try:
        extract_document_text("notes.txt", b"some text")
    except ValueError as error:
        assert "unsupported" in str(error).lower()
    else:
        raise AssertionError("Unsupported file type should be rejected")


def test_reject_empty_file():
    try:
        extract_document_text("empty.pdf", b"")
    except ValueError as error:
        assert "empty" in str(error).lower()
    else:
        raise AssertionError("Empty file should be rejected")


def test_reject_oversized_file():
    try:
        extract_document_text(
            "large.docx",
            b"x" * (10 * 1024 * 1024 + 1),
        )
    except ValueError as error:
        assert "10 mb" in str(error).lower()
    else:
        raise AssertionError("Oversized file should be rejected")


def test_agent_success_status():
    result = run_agent_with_status(
        "Grammar",
        lambda text: f"Reviewed: {text}",
        "Example",
    )

    assert result["agent"] == "Grammar"
    assert result["status"] == "success"
    assert result["output"] == "Reviewed: Example"
    assert result["duration"] >= 0
    assert result["error"] == ""


def test_agent_error_status():
    def failing_agent(_text):
        raise RuntimeError("Model unavailable")

    result = run_agent_with_status(
        "Fact Checker",
        failing_agent,
        "Example",
    )

    assert result["agent"] == "Fact Checker"
    assert result["status"] == "error"
    assert result["output"] == ""
    assert "Model unavailable" in result["error"]


def test_make_agent_result_rejects_invalid_status():
    try:
        make_agent_result("Grammar", "unknown", "", 0.1)
    except ValueError:
        pass
    else:
        raise AssertionError("Invalid status should be rejected")


def test_review_agent_captures_failure():
    from graph import document_review_workflow as review_module

    original_generate = review_module.generate

    try:
        def failing_generate(_prompt):
            raise RuntimeError("Simulated Gemini failure")

        review_module.generate = failing_generate

        result = review_module.research_review_node(
            {"document_text": "A sample document."}
        )

        agent_result = result["research_result"]

        assert agent_result["status"] == "error"
        assert agent_result["output"] == ""
        assert "Simulated Gemini failure" in agent_result["error"]

    finally:
        review_module.generate = original_generate


def test_review_agent_captures_success():
    from graph import document_review_workflow as review_module

    original_generate = review_module.generate

    try:
        review_module.generate = lambda _prompt: (
            "Review findings: add supporting evidence."
        )

        result = review_module.research_review_node(
            {"document_text": "A sample document."}
        )

        agent_result = result["research_result"]

        assert agent_result["status"] == "success"
        assert "supporting evidence" in agent_result["output"]

    finally:
        review_module.generate = original_generate


def test_review_document_rejects_empty_text():
    from graph.document_review_workflow import review_document

    try:
        review_document("   ")
    except ValueError as error:
        assert "empty" in str(error).lower()
    else:
        raise AssertionError("Empty document text should be rejected")


def test_file_size_boundary(self):
    from utils.document_review import (
        extract_document_text,
        MAX_FILE_SIZE,
    )

    # One byte over the configured maximum must be rejected.
    oversized = b"x" * (MAX_FILE_SIZE + 1)

    try:
        extract_document_text("sample.docx", oversized)
        raise AssertionError("Oversized upload should have been rejected.")
    except ValueError as error:
        assert "10 MB" in str(error)


def test_file_exactly_at_size_limit_passes_size_validation(self):
    from utils.document_review import (
        extract_document_text,
        MAX_FILE_SIZE,
    )

    # Use a valid DOCX padded with harmless ZIP comment bytes.
    from io import BytesIO
    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Boundary test document.")
    document.save(buffer)

    original = buffer.getvalue()

    # DOCX is a ZIP file; append bytes without modifying the ZIP contents.
    # The ZIP end-of-central-directory comment length is limited to 65535.
    padding_size = MAX_FILE_SIZE - len(original)
    if padding_size < 0 or padding_size > 65535:
        self.skipTest("Sample DOCX cannot be padded to the exact size limit.")

    # Create a valid archive comment of the required size.
    import struct

    data = bytearray(original)
    eocd = data.rfind(b"PK\x05\x06")
    if eocd < 0:
        self.fail("Could not locate ZIP end-of-central-directory record.")

    data[eocd + 20:eocd + 22] = struct.pack("<H", padding_size)
    data.extend(b"x" * padding_size)

    self.assertEqual(len(data), MAX_FILE_SIZE)

    text = extract_document_text("sample.docx", bytes(data))
    self.assertIn("Boundary test document", text)


if __name__ == "__main__":
    test_extract_docx_text()
    test_extract_pdf_text()
    test_reject_unsupported_file_type()
    test_reject_empty_file()
    test_reject_oversized_file()
    test_agent_success_status()
    test_agent_error_status()
    test_make_agent_result_rejects_invalid_status()
    test_review_agent_captures_failure()
    test_review_agent_captures_success()
    test_review_document_rejects_empty_text()

    print("Document-review tests passed successfully!")
