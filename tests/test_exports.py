
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader

from utils.export_pdf import create_pdf, create_document_review_pdf
from utils.export_docx import create_docx, create_document_review_docx


class TestReportExports(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_cwd = os.getcwd()
        os.chdir(self.temp_dir.name)
        os.makedirs("output", exist_ok=True)

    def tearDown(self):
        os.chdir(self.old_cwd)
        self.temp_dir.cleanup()

    def test_generated_report_pdf(self):
        path = create_pdf(
            "Test Topic",
            "This is a sample report.",
            "Review completed."
        )

        self.assertTrue(os.path.isfile(path))
        self.assertGreater(os.path.getsize(path), 0)

        reader = PdfReader(path)
        extracted = "\n".join(
            page.extract_text() or "" for page in reader.pages
        )

        self.assertIn("Test Topic", extracted)
        self.assertIn("sample report", extracted)

    def test_generated_report_docx(self):
        path = create_docx(
            "Test Topic",
            "This is a sample report.",
            "Review completed."
        )

        self.assertTrue(os.path.isfile(path))
        self.assertGreater(os.path.getsize(path), 0)

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)

        self.assertIn("Test Topic", text)
        self.assertIn("sample report", text)

    def test_review_pdf_export(self):
        results = [
            {
                "agent": "Grammar Agent",
                "status": "success",
                "output": "Correct the grammar in paragraph one.",
                "duration": 1.25,
                "error": "",
            }
        ]

        path = create_document_review_pdf(
            "sample.pdf",
            "sample.pdf",
            results
        )

        self.assertTrue(os.path.isfile(path))
        self.assertGreater(os.path.getsize(path), 0)

        reader = PdfReader(path)
        text = "\n".join(p.extract_text() or "" for p in reader.pages)

        self.assertIn("Grammar Agent", text)
        self.assertIn("Correct the grammar", text)

    def test_review_docx_export(self):
        results = [
            {
                "agent": "Reviewer Agent",
                "status": "success",
                "output": "Add evidence to the conclusion.",
                "duration": 0.75,
                "error": "",
            }
        ]

        path = create_document_review_docx(
            "sample.docx",
            "sample.docx",
            results
        )

        self.assertTrue(os.path.isfile(path))
        self.assertGreater(os.path.getsize(path), 0)

        document = Document(path)
        text = "\n".join(p.text for p in document.paragraphs)

        self.assertIn("Reviewer Agent", text)
        self.assertIn("Add evidence", text)


if __name__ == "__main__":
    unittest.main(verbosity=2)
