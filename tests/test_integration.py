
import os
import tempfile
import unittest
from unittest.mock import patch

from docx import Document
from pypdf import PdfReader, PdfWriter

from utils.document_review import extract_document_text
from utils.export_pdf import create_document_review_pdf
from utils.export_docx import create_document_review_docx
from graph.document_review_workflow import review_document


class TestDocumentReviewIntegration(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.old_cwd = os.getcwd()
        os.chdir(self.temp_dir.name)
        os.makedirs("output", exist_ok=True)

    def tearDown(self):
        os.chdir(self.old_cwd)
        self.temp_dir.cleanup()

    def make_docx(self):
        path = os.path.join(self.temp_dir.name, "sample.docx")
        document = Document()
        document.add_heading("Integration Test Document", level=1)
        document.add_paragraph(
            "This document contains text for integration testing."
        )
        document.save(path)
        return path

    def test_docx_extraction_then_review_exports(self):
        source_path = self.make_docx()

        with open(source_path, "rb") as file:
            extracted = extract_document_text(
                "sample.docx",
                file.read()
            )

        self.assertIn("integration testing", extracted)

        results = [
            {
                "agent": "Reviewer Agent",
                "status": "success",
                "output": "The document was extracted successfully.",
                "duration": 0.01,
                "error": "",
            }
        ]

        pdf_path = create_document_review_pdf(
            "sample.docx",
            "sample.docx",
            results
        )
        docx_path = create_document_review_docx(
            "sample.docx",
            "sample.docx",
            results
        )

        self.assertTrue(os.path.isfile(pdf_path))
        self.assertTrue(os.path.isfile(docx_path))

        pdf_text = "\n".join(
            page.extract_text() or ""
            for page in PdfReader(pdf_path).pages
        )
        self.assertIn("Reviewer Agent", pdf_text)

        exported_docx = Document(docx_path)
        exported_text = "\n".join(
            paragraph.text for paragraph in exported_docx.paragraphs
        )
        self.assertIn("Reviewer Agent", exported_text)

    def test_empty_document_is_rejected_before_review(self):
        with self.assertRaises(ValueError):
            review_document("   ")

    def test_review_workflow_with_mocked_llm(self):
        # Avoid real Gemini API calls; each agent gets deterministic output.
        with patch(
            "graph.document_review_workflow.generate",
            return_value="Mocked test finding."
        ):
            result = review_document(
                "A sample document with enough text to review."
            )

        for key in (
            "research_result",
            "outline_result",
            "writer_result",
            "grammar_result",
            "fact_checker_result",
            "citation_result",
            "reviewer_result",
        ):
            self.assertIn(key, result)
            self.assertEqual(result[key]["status"], "success")
            self.assertIn("Mocked test finding.", result[key]["output"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
