
import csv
import os
import platform
import statistics
import sys
import time

from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from io import BytesIO

from docx import Document
from pypdf import PdfReader, PdfWriter
# Ensure the project root is available when running this script directly.
PROJECT_ROOT = Path(__file__).resolve().parent.parent

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))
    
from utils.document_review import extract_document_text
from utils.export_pdf import create_document_review_pdf
from utils.export_docx import create_document_review_docx


PROJECT_ROOT = Path(__file__).resolve().parent.parent
RESULTS_DIR = PROJECT_ROOT / "test_results"
RESULTS_DIR.mkdir(parents=True, exist_ok=True)

CSV_PATH = RESULTS_DIR / "performance_results.csv"
ENV_PATH = RESULTS_DIR / "environment.txt"

RUNS = 5


def create_sample_docx(path):
    doc = Document()
    doc.add_heading("Performance Test Document", level=1)

    for i in range(100):
        doc.add_paragraph(
            f"Performance testing paragraph {i + 1}. "
            "This paragraph contains sample text for extraction timing."
        )

    doc.save(path)


def timed_runs(label, operation, runs=RUNS):
    durations = []

    for _ in range(runs):
        start = time.perf_counter()
        operation()
        durations.append(time.perf_counter() - start)

    return {
        "operation": label,
        "runs": runs,
        "mean_seconds": round(statistics.mean(durations), 6),
        "median_seconds": round(statistics.median(durations), 6),
        "min_seconds": round(min(durations), 6),
        "max_seconds": round(max(durations), 6),
        "status": "PASS",
        "timestamp": datetime.now().isoformat(timespec="seconds"),
    }


def main():
    results = []

    with TemporaryDirectory() as temp:
        temp_path = Path(temp)
        sample_docx = temp_path / "sample.docx"
        create_sample_docx(sample_docx)

        file_bytes = sample_docx.read_bytes()

        extracted_text = extract_document_text(
            sample_docx.name,
            file_bytes
        )

        if not extracted_text.strip():
            raise RuntimeError("Sample document extraction returned no text.")

        results.append(
            timed_runs(
                "DOCX text extraction",
                lambda: extract_document_text(
                    sample_docx.name,
                    file_bytes
                ),
            )
        )

        review_results = [
            {
                "agent": "Benchmark Reviewer",
                "status": "success",
                "output": "Synthetic benchmark finding.",
                "duration": 0.01,
                "error": "",
            }
        ]

        original_cwd = os.getcwd()

        try:
            os.chdir(temp_path)
            Path("output").mkdir(exist_ok=True)

            results.append(
                timed_runs(
                    "Review PDF export",
                    lambda: create_document_review_pdf(
                        "benchmark.docx",
                        "benchmark.docx",
                        review_results
                    ),
                )
            )

            results.append(
                timed_runs(
                    "Review DOCX export",
                    lambda: create_document_review_docx(
                        "benchmark.docx",
                        "benchmark.docx",
                        review_results
                    ),
                )
            )
        finally:
            os.chdir(original_cwd)

    with CSV_PATH.open("w", newline="", encoding="utf-8") as file:
        writer = csv.DictWriter(
            file,
            fieldnames=[
                "operation",
                "runs",
                "mean_seconds",
                "median_seconds",
                "min_seconds",
                "max_seconds",
                "status",
                "timestamp",
            ],
        )
        writer.writeheader()
        writer.writerows(results)

    environment = [
        f"Python: {sys.version}",
        f"Platform: {platform.platform()}",
        f"Processor: {platform.processor() or 'Not reported'}",
        f"Benchmark runs per operation: {RUNS}",
        f"Generated at: {datetime.now().isoformat(timespec='seconds')}",
        "Note: local extraction/export only; no Gemini API calls were made.",
    ]
    ENV_PATH.write_text("\n".join(environment) + "\n", encoding="utf-8")

    print("Performance benchmark completed.")
    print(f"CSV: {CSV_PATH}")
    print(f"Environment: {ENV_PATH}")

    for result in results:
        print(
            f"{result['operation']}: "
            f"mean={result['mean_seconds']}s, "
            f"median={result['median_seconds']}s "
            f"(n={result['runs']})"
        )


if __name__ == "__main__":
    main()
