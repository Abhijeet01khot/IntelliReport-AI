from docx import Document
from docx.shared import Pt
import os


def create_docx(topic, report, review):

    if not os.path.exists("output"):
        os.makedirs("output")

    filename = f"output/{topic.replace(' ','_')}.docx"

    doc = Document()

    title = doc.add_heading(topic, level=1)

    title.runs[0].font.size = Pt(20)

    doc.add_heading("Generated Report", level=2)

    doc.add_paragraph(report)

    doc.add_page_break()

    doc.add_heading("Reviewer Feedback", level=2)

    doc.add_paragraph(review)

    doc.save(filename)

    return filename


def create_document_review_docx(filename, original_name, review_results):
    """Export document-review findings to a Word document."""

    if not os.path.exists("output"):
        os.makedirs("output")

    safe_name = os.path.basename(filename)
    output_name = f"output/{os.path.splitext(safe_name)[0]}_Review.docx"

    doc = Document()
    doc.add_heading("IntelliReport AI - Document Review", level=1)
    doc.add_paragraph(f"Original document: {original_name}")

    for result in review_results:
        doc.add_heading(result.get("agent", "Agent"), level=2)
        doc.add_paragraph(
            f"Status: {result.get('status', 'unknown').upper()}"
        )
        doc.add_paragraph(
            f"Execution time: {result.get('duration', 0)} seconds"
        )

        if result.get("output"):
            doc.add_heading("Findings", level=3)
            doc.add_paragraph(result["output"])

        if result.get("error"):
            doc.add_heading("Error", level=3)
            doc.add_paragraph(result["error"])

    doc.save(output_name)
    return output_name
